from io import BytesIO

from flask import (
    Blueprint,
    render_template,
    request,
    send_file,
    flash,
    redirect,
    url_for
)

from flask_login import login_required

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)

from models.group import ChitGroup
from models.customer import Customer

from services.report_service import (
    get_group_report,
    get_payment_report,
    get_auction_report
)


reports = Blueprint(
    "reports",
    __name__,
    url_prefix="/reports"
)


# ============================================================
# HELPERS
# ============================================================

def get_report_parameters():
    """
    Read common report filters from the request.

    Supported:
        today
        weekly
        monthly
        yearly
        custom

    Multiple group IDs are supported.
    """

    period = request.args.get(
        "period",
        "monthly"
    )

    group_ids = request.args.getlist(
        "group_ids"
    )

    from_date = request.args.get(
        "from_date",
        ""
    )

    to_date = request.args.get(
        "to_date",
        ""
    )

    selected_group_ids = []

    for group_id in group_ids:

        if not group_id:
            continue

        try:

            selected_group_ids.append(
                int(group_id)
            )

        except ValueError:

            continue

    return (
        period,
        selected_group_ids,
        from_date,
        to_date
    )


def get_payment_report_parameters():
    """
    Read payment report filters.
    """

    period = request.args.get(
        "period",
        "monthly"
    )

    group_ids = request.args.getlist(
        "group_ids"
    )

    customer_id = request.args.get(
        "customer_id",
        ""
    )

    payment_method = request.args.get(
        "payment_method",
        ""
    )

    from_date = request.args.get(
        "from_date",
        ""
    )

    to_date = request.args.get(
        "to_date",
        ""
    )

    selected_group_ids = []

    for group_id in group_ids:

        if not group_id:
            continue

        try:

            selected_group_ids.append(
                int(group_id)
            )

        except ValueError:

            continue

    return (
        period,
        selected_group_ids,
        customer_id,
        payment_method,
        from_date,
        to_date
    )


# ============================================================
# GROUP REPORT BUILDER
# ============================================================

def build_group_report():

    (
        period,
        group_ids,
        from_date,
        to_date
    ) = get_report_parameters()

    report = get_group_report(
        period=period,
        from_date=from_date or None,
        to_date=to_date or None,
        group_ids=group_ids
    )

    return (
        report,
        period,
        group_ids,
        from_date,
        to_date
    )


# ============================================================
# PAYMENT REPORT BUILDER
# ============================================================

def build_payment_report():

    (
        period,
        group_ids,
        customer_id,
        payment_method,
        from_date,
        to_date
    ) = get_payment_report_parameters()

    report = get_payment_report(
        period=period,
        from_date=from_date or None,
        to_date=to_date or None,
        group_ids=group_ids,
        customer_id=customer_id or None,
        payment_method=payment_method or None
    )

    return (
        report,
        period,
        group_ids,
        customer_id,
        payment_method,
        from_date,
        to_date
    )


# ============================================================
# AUCTION REPORT BUILDER
# ============================================================

def build_auction_report():

    (
        period,
        group_ids,
        from_date,
        to_date
    ) = get_report_parameters()

    report = get_auction_report(
        period=period,
        from_date=from_date or None,
        to_date=to_date or None,
        group_ids=group_ids
    )

    return (
        report,
        period,
        group_ids,
        from_date,
        to_date
    )


# ============================================================
# REPORT CENTER
# ============================================================

@reports.route("/")
@login_required
def index():

    groups = ChitGroup.query.order_by(
        ChitGroup.name
    ).all()

    customers = Customer.query.filter_by(
        status="ACTIVE"
    ).order_by(
        Customer.name
    ).all()

    return render_template(
        "reports/index.html",
        groups=groups,
        customers=customers
    )


# ============================================================
# GROUP REPORT
# ============================================================

@reports.route(
    "/group",
    methods=["GET"]
)
@login_required
def group_report():

    try:

        (
            report,
            period,
            group_ids,
            from_date,
            to_date
        ) = build_group_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    return render_template(
        "reports/group_report.html",

        report=report,

        period=period,

        selected_group_ids=[
            str(group_id)
            for group_id in group_ids
        ],

        from_date=from_date,

        to_date=to_date
    )


# ============================================================
# GROUP REPORT PDF
# ============================================================

@reports.route(
    "/group/pdf",
    methods=["GET"]
)
@login_required
def group_report_pdf():

    try:

        (
            report,
            period,
            group_ids,
            from_date,
            to_date
        ) = build_group_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    buffer = BytesIO()

    document = SimpleDocTemplate(

        buffer,

        pagesize=landscape(A4),

        rightMargin=10 * mm,

        leftMargin=10 * mm,

        topMargin=10 * mm,

        bottomMargin=10 * mm
    )

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "ABR CHIT FUND",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            "Group Installment Report",
            styles["Heading2"]
        )
    )

    if report["start_date"] and report["end_date"]:

        date_text = (
            f"{report['start_date'].strftime('%d-%m-%Y')}"
            f" to "
            f"{report['end_date'].strftime('%d-%m-%Y')}"
        )

    else:

        date_text = "All Dates"

    elements.append(
        Paragraph(
            f"Period: {date_text}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            8
        )
    )

    if report["groups"]:

        selected_groups_text = ", ".join(
            group.name
            for group in report["groups"]
        )

    else:

        selected_groups_text = "All Groups"

    elements.append(
        Paragraph(
            f"Groups: {selected_groups_text}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            8
        )
    )

    summary = report["summary"]

    summary_data = [

        [
            "Groups",
            "Members",
            "Installments",
            "Total Due",
            "Total Paid",
            "Pending"
        ],

        [
            str(summary["groups"]),

            str(summary["members"]),

            str(summary["installments"]),

            f"₹{summary['total_due']:,.2f}",

            f"₹{summary['total_paid']:,.2f}",

            f"₹{summary['total_pending']:,.2f}"
        ]
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            30 * mm,
            30 * mm,
            35 * mm,
            40 * mm,
            40 * mm,
            40 * mm
        ]
    )

    summary_table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1f4e78")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "PADDING",
                (0, 0),
                (-1, -1),
                6
            )
        ])
    )

    elements.append(
        summary_table
    )

    elements.append(
        Spacer(
            1,
            10
        )
    )

    data = [

        [
            "Group",
            "Customer",
            "Mobile",
            "Ticket",
            "Installment",
            "Due Month",
            "Due",
            "Paid",
            "Pending",
            "Status"
        ]
    ]

    for row in report["rows"]:

        due_month = row["due_month"]

        if due_month:

            due_month_text = due_month.strftime(
                "%b-%Y"
            )

        else:

            due_month_text = ""

        data.append([

            row["group_name"],

            row["customer_name"],

            row["customer_mobile"],

            str(
                row["ticket_number"]
            ),

            str(
                row["installment_number"]
            ),

            due_month_text,

            f"₹{row['due_amount']:,.2f}",

            f"₹{row['paid_amount']:,.2f}",

            f"₹{row['balance_amount']:,.2f}",

            row["status"]
        ])

    if len(data) == 1:

        data.append([
            "No records",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            ""
        ])

    table = Table(
        data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#343a40")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.3,
                colors.grey
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),

            (
                "ALIGN",
                (6, 1),
                (8, -1),
                "RIGHT"
            ),

            (
                "PADDING",
                (0, 0),
                (-1, -1),
                4
            )
        ])
    )

    elements.append(
        table
    )

    document.build(
        elements
    )

    buffer.seek(0)

    return send_file(

        buffer,

        as_attachment=True,

        download_name="ABR_Group_Report.pdf",

        mimetype="application/pdf"
    )


# ============================================================
# GROUP REPORT EXCEL
# ============================================================

@reports.route(
    "/group/excel",
    methods=["GET"]
)
@login_required
def group_report_excel():

    try:

        (
            report,
            period,
            group_ids,
            from_date,
            to_date
        ) = build_group_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Group Report"

    worksheet["A1"] = "ABR CHIT FUND"

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A2"] = "Group Installment Report"

    worksheet["A2"].font = Font(
        bold=True,
        size=13
    )

    if report["start_date"] and report["end_date"]:

        worksheet["A3"] = (
            f"Period: "
            f"{report['start_date'].strftime('%d-%m-%Y')}"
            f" to "
            f"{report['end_date'].strftime('%d-%m-%Y')}"
        )

    else:

        worksheet["A3"] = "Period: All Dates"

    summary = report["summary"]

    worksheet.append([])

    worksheet.append([

        "Groups",
        "Members",
        "Installments",
        "Total Due",
        "Total Paid",
        "Pending"
    ])

    worksheet.append([

        summary["groups"],
        summary["members"],
        summary["installments"],
        summary["total_due"],
        summary["total_paid"],
        summary["total_pending"]
    ])

    worksheet.append([])

    headers = [

        "Group",
        "Customer",
        "Mobile",
        "Ticket",
        "Installment",
        "Due Month",
        "Due Amount",
        "Paid Amount",
        "Pending Amount",
        "Status"
    ]

    worksheet.append(
        headers
    )

    for cell in worksheet[
        worksheet.max_row
    ]:

        cell.font = Font(
            bold=True
        )

    for row in report["rows"]:

        due_month = row["due_month"]

        due_month_text = ""

        if due_month:

            due_month_text = due_month.strftime(
                "%d-%m-%Y"
            )

        worksheet.append([

            row["group_name"],
            row["customer_name"],
            row["customer_mobile"],
            row["ticket_number"],
            row["installment_number"],
            due_month_text,
            row["due_amount"],
            row["paid_amount"],
            row["balance_amount"],
            row["status"]
        ])

    if not report["rows"]:

        worksheet.append([
            "No records found"
        ])

    widths = {

        "A": 25,
        "B": 25,
        "C": 18,
        "D": 10,
        "E": 15,
        "F": 15,
        "G": 15,
        "H": 15,
        "I": 15,
        "J": 15
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="center"
            )

    buffer = BytesIO()

    workbook.save(
        buffer
    )

    buffer.seek(0)

    return send_file(

        buffer,

        as_attachment=True,

        download_name="ABR_Group_Report.xlsx",

        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )


# ============================================================
# PAYMENT REPORT
# ============================================================

@reports.route(
    "/payment",
    methods=["GET"]
)
@login_required
def payment_report():

    try:

        (
            report,
            period,
            group_ids,
            customer_id,
            payment_method,
            from_date,
            to_date
        ) = build_payment_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    customers = Customer.query.filter_by(
        status="ACTIVE"
    ).order_by(
        Customer.name
    ).all()

    groups = ChitGroup.query.order_by(
        ChitGroup.name
    ).all()

    return render_template(
        "reports/payment_report.html",

        report=report,

        period=period,

        selected_group_ids=[
            str(group_id)
            for group_id in group_ids
        ],

        selected_customer_id=(
            str(customer_id)
            if customer_id
            else ""
        ),

        selected_payment_method=(
            payment_method
            or ""
        ),

        from_date=from_date,

        to_date=to_date,

        customers=customers,

        groups=groups
    )


# ============================================================
# PAYMENT REPORT PDF
# ============================================================

@reports.route(
    "/payment/pdf",
    methods=["GET"]
)
@login_required
def payment_report_pdf():

    try:

        (
            report,
            period,
            group_ids,
            customer_id,
            payment_method,
            from_date,
            to_date
        ) = build_payment_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    buffer = BytesIO()

    document = SimpleDocTemplate(

        buffer,

        pagesize=landscape(A4),

        rightMargin=10 * mm,

        leftMargin=10 * mm,

        topMargin=10 * mm,

        bottomMargin=10 * mm
    )

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "ABR CHIT FUND",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            "Payment Collection Report",
            styles["Heading2"]
        )
    )

    if report["start_date"] and report["end_date"]:

        date_text = (
            f"{report['start_date'].strftime('%d-%m-%Y')}"
            f" to "
            f"{report['end_date'].strftime('%d-%m-%Y')}"
        )

    else:

        date_text = "All Dates"

    elements.append(
        Paragraph(
            f"Period: {date_text}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            8
        )
    )

    summary = report["summary"]

    summary_data = [

        [
            "Payments",
            "Customers",
            "Groups",
            "Total Received"
        ],

        [
            str(summary["payments"]),

            str(summary["customers"]),

            str(summary["groups"]),

            f"₹{summary['total_received']:,.2f}"
        ]
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            40 * mm,
            40 * mm,
            40 * mm,
            50 * mm
        ]
    )

    summary_table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#198754")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "PADDING",
                (0, 0),
                (-1, -1),
                6
            )
        ])
    )

    elements.append(
        summary_table
    )

    elements.append(
        Spacer(
            1,
            10
        )
    )

    data = [

        [
            "Receipt",
            "Date",
            "Customer",
            "Mobile",
            "Groups",
            "Amount",
            "Method",
            "Reference"
        ]
    ]

    for row in report["rows"]:

        data.append([

            row["receipt_number"],

            row["payment_date"].strftime(
                "%d-%m-%Y"
            ),

            row["customer_name"],

            row["customer_mobile"],

            row["group_text"],

            f"₹{row['amount']:,.2f}",

            row["payment_method"],

            row["transaction_reference"]
        ])

    if len(data) == 1:

        data.append([
            "No records",
            "",
            "",
            "",
            "",
            "",
            "",
            ""
        ])

    table = Table(
        data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#343a40")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.3,
                colors.grey
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),

            (
                "ALIGN",
                (5, 1),
                (5, -1),
                "RIGHT"
            ),

            (
                "PADDING",
                (0, 0),
                (-1, -1),
                4
            )
        ])
    )

    elements.append(
        table
    )

    document.build(
        elements
    )

    buffer.seek(0)

    return send_file(

        buffer,

        as_attachment=True,

        download_name="ABR_Payment_Report.pdf",

        mimetype="application/pdf"
    )


# ============================================================
# PAYMENT REPORT EXCEL
# ============================================================

@reports.route(
    "/payment/excel",
    methods=["GET"]
)
@login_required
def payment_report_excel():

    try:

        (
            report,
            period,
            group_ids,
            customer_id,
            payment_method,
            from_date,
            to_date
        ) = build_payment_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Payment Report"

    worksheet["A1"] = "ABR CHIT FUND"

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A2"] = "Payment Collection Report"

    worksheet["A2"].font = Font(
        bold=True,
        size=13
    )

    if report["start_date"] and report["end_date"]:

        worksheet["A3"] = (
            f"Period: "
            f"{report['start_date'].strftime('%d-%m-%Y')}"
            f" to "
            f"{report['end_date'].strftime('%d-%m-%Y')}"
        )

    else:

        worksheet["A3"] = "Period: All Dates"

    summary = report["summary"]

    worksheet.append([])

    worksheet.append([

        "Payments",
        "Customers",
        "Groups",
        "Total Received"
    ])

    worksheet.append([

        summary["payments"],
        summary["customers"],
        summary["groups"],
        summary["total_received"]
    ])

    worksheet.append([])

    headers = [

        "Receipt Number",
        "Payment Date",
        "Customer",
        "Mobile",
        "Groups",
        "Amount",
        "Payment Method",
        "Transaction Reference",
        "Notes"
    ]

    worksheet.append(
        headers
    )

    for cell in worksheet[
        worksheet.max_row
    ]:

        cell.font = Font(
            bold=True
        )

    for row in report["rows"]:

        worksheet.append([

            row["receipt_number"],

            row["payment_date"].strftime(
                "%d-%m-%Y"
            ),

            row["customer_name"],

            row["customer_mobile"],

            row["group_text"],

            row["amount"],

            row["payment_method"],

            row["transaction_reference"],

            row["notes"]
        ])

    if not report["rows"]:

        worksheet.append([
            "No records found"
        ])

    widths = {

        "A": 25,
        "B": 15,
        "C": 25,
        "D": 18,
        "E": 30,
        "F": 18,
        "G": 18,
        "H": 25,
        "I": 30
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="center"
            )

    buffer = BytesIO()

    workbook.save(
        buffer
    )

    buffer.seek(0)

    return send_file(

        buffer,

        as_attachment=True,

        download_name="ABR_Payment_Report.xlsx",

        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )


# ============================================================
# AUCTION REPORT
# ============================================================

@reports.route(
    "/auction",
    methods=["GET"]
)
@login_required
def auction_report():

    try:

        (
            report,
            period,
            group_ids,
            from_date,
            to_date
        ) = build_auction_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    groups = ChitGroup.query.order_by(
        ChitGroup.name
    ).all()

    return render_template(
        "reports/auction_report.html",

        report=report,

        period=period,

        selected_group_ids=[
            str(group_id)
            for group_id in group_ids
        ],

        from_date=from_date,

        to_date=to_date,

        groups=groups
    )


# ============================================================
# AUCTION REPORT PDF
# ============================================================

@reports.route(
    "/auction/pdf",
    methods=["GET"]
)
@login_required
def auction_report_pdf():

    try:

        (
            report,
            period,
            group_ids,
            from_date,
            to_date
        ) = build_auction_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    buffer = BytesIO()

    document = SimpleDocTemplate(

        buffer,

        pagesize=landscape(A4),

        rightMargin=10 * mm,

        leftMargin=10 * mm,

        topMargin=10 * mm,

        bottomMargin=10 * mm
    )

    styles = getSampleStyleSheet()

    elements = []

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    elements.append(
        Paragraph(
            "ABR CHIT FUND",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            "Auction Report",
            styles["Heading2"]
        )
    )

    # --------------------------------------------------------
    # Date Range
    # --------------------------------------------------------

    if (
        report["start_date"]
        and report["end_date"]
    ):

        date_text = (
            f"{report['start_date'].strftime('%d-%m-%Y')}"
            f" to "
            f"{report['end_date'].strftime('%d-%m-%Y')}"
        )

    else:

        date_text = "All Dates"

    elements.append(
        Paragraph(
            f"Period: {date_text}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            8
        )
    )

    # --------------------------------------------------------
    # Groups
    # --------------------------------------------------------

    if group_ids:

        selected_groups = ChitGroup.query.filter(
            ChitGroup.id.in_(group_ids)
        ).order_by(
            ChitGroup.name
        ).all()

        selected_groups_text = ", ".join(
            group.name
            for group in selected_groups
        )

    else:

        selected_groups_text = "All Groups"

    elements.append(
        Paragraph(
            f"Groups: {selected_groups_text}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            8
        )
    )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    summary = report["summary"]

    summary_data = [

        [
            "Auctions",
            "Groups",
            "Winners",
            "Total Bid",
            "Commission",
            "Net Bid",
            "Installment"
        ],

        [

            str(
                summary["auctions"]
            ),

            str(
                summary["groups"]
            ),

            str(
                summary["winners"]
            ),

            f"₹{summary['total_bid_amount']:,.2f}",

            f"₹{summary['total_commission']:,.2f}",

            f"₹{summary['total_net_bid']:,.2f}",

            f"₹{summary['total_installment']:,.2f}"
        ]
    ]

    summary_table = Table(
        summary_data,

        colWidths=[
            25 * mm,
            25 * mm,
            25 * mm,
            35 * mm,
            35 * mm,
            35 * mm,
            35 * mm
        ]
    )

    summary_table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#6f42c1")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "PADDING",
                (0, 0),
                (-1, -1),
                6
            )
        ])
    )

    elements.append(
        summary_table
    )

    elements.append(
        Spacer(
            1,
            10
        )
    )

    # --------------------------------------------------------
    # Auction Table
    # --------------------------------------------------------

    data = [

        [
            "Date",
            "Month",
            "Group",
            "Winner",
            "Mobile",
            "Bid Amount",
            "Commission",
            "Net Bid",
            "Installment",
            "Status"
        ]
    ]

    for row in report["rows"]:

        auction_date = row["auction_date"]

        if auction_date:

            auction_date_text = (
                auction_date.strftime(
                    "%d-%m-%Y"
                )
            )

        else:

            auction_date_text = ""

        data.append([

            auction_date_text,

            str(
                row["auction_month"]
            ),

            row["group_name"],

            row["winner_name"],

            row["winner_mobile"],

            f"₹{row['bid_amount']:,.2f}",

            f"₹{row['commission_amount']:,.2f}",

            f"₹{row['net_bid_amount']:,.2f}",

            f"₹{row['installment_amount']:,.2f}",

            row["status"]
        ])

    if len(data) == 1:

        data.append([
            "No records",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            ""
        ])

    table = Table(
        data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#343a40")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.3,
                colors.grey
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),

            (
                "ALIGN",
                (5, 1),
                (8, -1),
                "RIGHT"
            ),

            (
                "PADDING",
                (0, 0),
                (-1, -1),
                4
            )
        ])
    )

    elements.append(
        table
    )

    document.build(
        elements
    )

    buffer.seek(0)

    return send_file(

        buffer,

        as_attachment=True,

        download_name="ABR_Auction_Report.pdf",

        mimetype="application/pdf"
    )


# ============================================================
# AUCTION REPORT EXCEL
# ============================================================

@reports.route(
    "/auction/excel",
    methods=["GET"]
)
@login_required
def auction_report_excel():

    try:

        (
            report,
            period,
            group_ids,
            from_date,
            to_date
        ) = build_auction_report()

    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return redirect(
            url_for(
                "reports.index"
            )
        )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Auction Report"

    # --------------------------------------------------------
    # Header
    # --------------------------------------------------------

    worksheet["A1"] = "ABR CHIT FUND"

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A2"] = "Auction Report"

    worksheet["A2"].font = Font(
        bold=True,
        size=13
    )

    # --------------------------------------------------------
    # Period
    # --------------------------------------------------------

    if (
        report["start_date"]
        and report["end_date"]
    ):

        worksheet["A3"] = (
            f"Period: "
            f"{report['start_date'].strftime('%d-%m-%Y')}"
            f" to "
            f"{report['end_date'].strftime('%d-%m-%Y')}"
        )

    else:

        worksheet["A3"] = (
            "Period: All Dates"
        )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    summary = report["summary"]

    worksheet.append([])

    worksheet.append([

        "Auctions",
        "Groups",
        "Winners",
        "Total Bid",
        "Commission",
        "Net Bid",
        "Installment"
    ])

    worksheet.append([

        summary["auctions"],

        summary["groups"],

        summary["winners"],

        summary["total_bid_amount"],

        summary["total_commission"],

        summary["total_net_bid"],

        summary["total_installment"]
    ])

    worksheet.append([])

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [

        "Auction Date",
        "Auction Month",
        "Group",
        "Winner",
        "Mobile",
        "Bid Amount",
        "Commission Amount",
        "Net Bid Amount",
        "Installment Amount",
        "Status"
    ]

    worksheet.append(
        headers
    )

    for cell in worksheet[
        worksheet.max_row
    ]:

        cell.font = Font(
            bold=True
        )

    # --------------------------------------------------------
    # Rows
    # --------------------------------------------------------

    for row in report["rows"]:

        auction_date = row["auction_date"]

        if auction_date:

            auction_date_text = (
                auction_date.strftime(
                    "%d-%m-%Y"
                )
            )

        else:

            auction_date_text = ""

        worksheet.append([

            auction_date_text,

            row["auction_month"],

            row["group_name"],

            row["winner_name"],

            row["winner_mobile"],

            row["bid_amount"],

            row["commission_amount"],

            row["net_bid_amount"],

            row["installment_amount"],

            row["status"]
        ])

    if not report["rows"]:

        worksheet.append([
            "No records found"
        ])

    # --------------------------------------------------------
    # Column Widths
    # --------------------------------------------------------

    widths = {

        "A": 18,
        "B": 18,
        "C": 25,
        "D": 25,
        "E": 18,
        "F": 18,
        "G": 20,
        "H": 20,
        "I": 20,
        "J": 15
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # Alignment
    # --------------------------------------------------------

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="center"
            )

    # --------------------------------------------------------
    # Create File
    # --------------------------------------------------------

    buffer = BytesIO()

    workbook.save(
        buffer
    )

    buffer.seek(0)

    return send_file(

        buffer,

        as_attachment=True,

        download_name="ABR_Auction_Report.xlsx",

        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )